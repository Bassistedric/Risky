from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook


# ============================================================
# IMPORT RH — MAPPING VMA
# ============================================================

RH_ORGANIZATION_MAP = {
    "V.M.A.": "VMA_NORD",
    "VMA BE.MAINTENANCE SA": "VMA_MAINTENANCE",
    "VMA SUD": "VMA_SUD",
}

WORKFORCE_MAP = {
    "A": "WORKER",
    "B": "EMPLOYEE",
}


# ============================================================
# UTILITAIRES
# ============================================================

def _as_float(value: Any) -> float:
    if value in (None, ""):
        return 0.0

    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Valeur d'heures invalide : {value!r}"
        ) from exc


def _find_hours_sheet(workbook):
    exact = [
        name
        for name in workbook.sheetnames
        if re.fullmatch(
            r"\d{4}\s+Gewerkte uren",
            name.strip(),
            flags=re.IGNORECASE,
        )
    ]

    if exact:
        return workbook[exact[0]]

    candidates = [
        name
        for name in workbook.sheetnames
        if "gewerkte uren" in name.lower()
    ]

    if not candidates:
        raise ValueError(
            "Aucune feuille 'Gewerkte uren' trouvée"
        )

    return workbook[candidates[0]]


def _read_year_and_month_columns(sheet):
    month_columns: list[tuple[int, int]] = []
    years: set[int] = set()

    for column in range(
        2,
        sheet.max_column + 1,
    ):
        header = sheet.cell(
            row=1,
            column=column,
        ).value

        if not header:
            continue

        match = re.search(
            r"(20\d{2})",
            str(header),
        )

        if not match:
            continue

        year = int(match.group(1))
        years.add(year)

        month_columns.append(
            (
                column,
                len(month_columns) + 1,
            )
        )

        if len(month_columns) == 12:
            break

    if len(years) != 1:
        raise ValueError(
            "Impossible de déterminer une année unique "
            "dans les colonnes d'heures"
        )

    if not month_columns:
        raise ValueError(
            "Aucune colonne mensuelle trouvée"
        )

    return years.pop(), month_columns


def _find_last_available_month(
    sheet,
    month_columns: list[tuple[int, int]],
) -> int:
    grand_total_row = None

    for row in range(
        1,
        sheet.max_row + 1,
    ):
        label = sheet.cell(
            row=row,
            column=1,
        ).value

        if (
            isinstance(label, str)
            and label.strip().lower()
            == "eindtotaal"
        ):
            grand_total_row = row
            break

    source_rows = (
        [grand_total_row]
        if grand_total_row
        else list(
            range(
                1,
                sheet.max_row + 1,
            )
        )
    )

    last_month = 0

    for column, month in month_columns:
        has_value = any(
            abs(
                _as_float(
                    sheet.cell(
                        row=row,
                        column=column,
                    ).value
                )
            )
            > 0
            for row in source_rows
        )

        if has_value:
            last_month = month

    if last_month == 0:
        raise ValueError(
            "Aucune heure prestée détectée "
            "dans le fichier"
        )

    return last_month


# ============================================================
# PARSING DU FICHIER RH
# ============================================================

def parse_rh_work_hours(
    file_bytes: bytes,
    filename: str,
) -> dict:
    if not file_bytes:
        raise ValueError(
            "Le fichier transmis est vide"
        )

    try:
        workbook = load_workbook(
            BytesIO(file_bytes),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Le fichier Excel ne peut pas être lu"
        ) from exc

    sheet = _find_hours_sheet(
        workbook
    )

    year, month_columns = (
        _read_year_and_month_columns(
            sheet
        )
    )

    month_to = _find_last_available_month(
        sheet,
        month_columns,
    )

    extracted: dict[
        str,
        dict[str, list[float]],
    ] = {}

    current_source_org: str | None = None

    for row in range(
        2,
        sheet.max_row + 1,
    ):
        raw_label = sheet.cell(
            row=row,
            column=1,
        ).value

        if raw_label is None:
            continue

        label = str(raw_label).strip()

        if label in RH_ORGANIZATION_MAP:
            current_source_org = label
            extracted.setdefault(
                label,
                {},
            )
            continue

        if (
            current_source_org
            and label in WORKFORCE_MAP
        ):
            category = WORKFORCE_MAP[
                label
            ]

            values = [
                _as_float(
                    sheet.cell(
                        row=row,
                        column=column,
                    ).value
                )
                for column, month
                in month_columns
                if month <= month_to
            ]

            extracted[
                current_source_org
            ][category] = values

    missing = []

    for source_org in RH_ORGANIZATION_MAP:
        categories = extracted.get(
            source_org,
            {},
        )

        for category in (
            "WORKER",
            "EMPLOYEE",
        ):
            if category not in categories:
                missing.append(
                    f"{source_org} / {category}"
                )

    if missing:
        raise ValueError(
            "Données RH incomplètes : "
            + ", ".join(missing)
        )

    rows: list[dict] = []
    entities: list[dict] = []

    for (
        source_org,
        organization_code,
    ) in RH_ORGANIZATION_MAP.items():
        worker_values = extracted[
            source_org
        ]["WORKER"]

        employee_values = extracted[
            source_org
        ]["EMPLOYEE"]

        entity_rows = []

        for month in range(
            1,
            month_to + 1,
        ):
            worker_hours = (
                worker_values[
                    month - 1
                ]
            )

            employee_hours = (
                employee_values[
                    month - 1
                ]
            )

            row = {
                "organization_code":
                    organization_code,
                "source_label":
                    source_org,
                "year": year,
                "month": month,
                "worker_hours":
                    worker_hours,
                "employee_hours":
                    employee_hours,
                "total_hours":
                    worker_hours
                    + employee_hours,
            }

            rows.append(row)
            entity_rows.append(row)

        entities.append(
            {
                "organization_code":
                    organization_code,
                "source_label":
                    source_org,
                "year": year,
                "month_to": month_to,
                "worker_hours_ytd": sum(
                    row["worker_hours"]
                    for row in entity_rows
                ),
                "employee_hours_ytd": sum(
                    row["employee_hours"]
                    for row in entity_rows
                ),
                "total_hours_ytd": sum(
                    row["total_hours"]
                    for row in entity_rows
                ),
                "latest_month":
                    entity_rows[-1],
            }
        )

    return {
        "filename": filename,
        "sheet_name": sheet.title,
        "year": year,
        "month_to": month_to,
        "entities": entities,
        "rows": rows,
        "record_count": (
            len(rows) * 2
        ),
    }
